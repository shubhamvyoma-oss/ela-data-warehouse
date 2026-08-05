from __future__ import annotations

import argparse
import csv
import hashlib
import re
import uuid
from collections.abc import Iterator
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from psycopg2.extras import Json, execute_values

from api_scripts.common.models import payload_sha256
from api_scripts.common.repositories import RunRepository
from shared.config import DatabaseSettings, WarehouseSettings
from shared.database import Database
from shared.logging import configure_logging

SOURCE_PATTERN = re.compile(r"^[a-z][a-z0-9_]{1,62}$")
SUPPORTED_SUFFIXES = {".csv", ".tsv", ".xlsx"}


class ImportValidationError(ValueError):
    pass


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def iter_rows(
    path: Path, sheet_name: str | None = None
) -> tuple[list[str], Iterator[tuple[int, dict[str, Any]]]]:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ImportValidationError(f"unsupported file type {suffix!r}")
    if suffix in {".csv", ".tsv"}:
        return _csv_rows(path, "\t" if suffix == ".tsv" else ",")
    return _xlsx_rows(path, sheet_name)


def _validate_headers(values: list[Any]) -> list[str]:
    headers = [str(value).strip() if value is not None else "" for value in values]
    if not headers or any(not header for header in headers):
        raise ImportValidationError("all source columns must have non-empty headers")
    normalized = [header.casefold() for header in headers]
    if len(set(normalized)) != len(normalized):
        raise ImportValidationError("source headers must be unique ignoring case")
    return headers


def _csv_rows(path: Path, delimiter: str) -> tuple[list[str], Iterator[tuple[int, dict[str, Any]]]]:
    stream = path.open("r", newline="", encoding="utf-8-sig")
    reader = csv.reader(stream, delimiter=delimiter)
    try:
        headers = _validate_headers(next(reader))
    except StopIteration as exc:
        stream.close()
        raise ImportValidationError("source file is empty") from exc

    def rows() -> Iterator[tuple[int, dict[str, Any]]]:
        try:
            for row_number, values in enumerate(reader, start=2):
                if not any(value.strip() for value in values):
                    continue
                if len(values) != len(headers):
                    raise ImportValidationError(
                        f"row {row_number} has {len(values)} values; expected {len(headers)}"
                    )
                yield row_number, dict(zip(headers, values, strict=True))
        finally:
            stream.close()

    return headers, rows()


def _xlsx_rows(path: Path, sheet_name: str | None) -> tuple[list[str], Iterator[tuple[int, dict[str, Any]]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            workbook.close()
            raise ImportValidationError(f"worksheet {sheet_name!r} does not exist")
        worksheet = workbook[sheet_name]
    else:
        worksheet = workbook.active
    row_iterator = worksheet.iter_rows(values_only=True)
    try:
        headers = _validate_headers(list(next(row_iterator)))
    except StopIteration as exc:
        workbook.close()
        raise ImportValidationError("worksheet is empty") from exc

    def rows() -> Iterator[tuple[int, dict[str, Any]]]:
        try:
            for row_number, values in enumerate(row_iterator, start=2):
                if not any(value not in (None, "") for value in values):
                    continue
                if len(values) != len(headers):
                    raise ImportValidationError(
                        f"row {row_number} has {len(values)} values; expected {len(headers)}"
                    )
                yield row_number, dict(zip(headers, values, strict=True))
        finally:
            workbook.close()

    return headers, rows()


def import_file(
    database: Database,
    *,
    source_name: str,
    path: Path,
    required_columns: tuple[str, ...] = (),
    sheet_name: str | None = None,
    batch_size: int = 500,
) -> uuid.UUID:
    if not SOURCE_PATTERN.fullmatch(source_name):
        raise ImportValidationError("source name must use lowercase letters, numbers, and underscores")
    path = path.resolve(strict=True)
    headers, rows = iter_rows(path, sheet_name)
    missing = sorted(set(required_columns) - set(headers))
    if missing:
        raise ImportValidationError(f"required columns are missing: {', '.join(missing)}")

    checksum = file_sha256(path)
    runs = RunRepository(database)
    pipeline_name = f"manual_import:{source_name}"
    run_id = runs.start(pipeline_name, "manual", {})
    import_id = uuid.uuid4()
    rows_read = 0
    rows_written = 0
    try:
        with database.transaction() as connection, connection.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO audit.manual_imports (
                    id, run_id, source_name, original_filename, file_sha256,
                    file_size_bytes, status
                ) VALUES (%s, %s, %s, %s, %s, %s, 'VALIDATING')
                """,
                (import_id, run_id, source_name, path.name, checksum, path.stat().st_size),
            )

        batch: list[tuple[Any, ...]] = []
        for row_number, row in rows:
            rows_read += 1
            batch.append(
                (
                    import_id,
                    run_id,
                    source_name,
                    row_number,
                    payload_sha256(row),
                    Json(row),
                )
            )
            if len(batch) >= batch_size:
                rows_written += _write_batch(database, batch)
                batch.clear()
        if batch:
            rows_written += _write_batch(database, batch)

        with database.transaction() as connection, connection.cursor() as cursor:
            cursor.execute(
                """
                UPDATE audit.manual_imports
                SET status = 'LOADED', rows_read = %s, rows_loaded = %s
                WHERE id = %s
                """,
                (rows_read, rows_written, import_id),
            )
        runs.finish(
            run_id,
            pipeline_name,
            status="succeeded",
            rows_read=rows_read,
            rows_written=rows_written,
            rows_rejected=0,
            checkpoint_after={"file_sha256": checksum, "import_id": str(import_id)},
            request_count=0,
        )
        return import_id
    except Exception as exc:
        with database.transaction() as connection, connection.cursor() as cursor:
            cursor.execute(
                """
                UPDATE audit.manual_imports
                SET status = 'FAILED', rows_read = %s, rows_loaded = %s
                WHERE id = %s
                """,
                (rows_read, rows_written, import_id),
            )
        runs.finish(
            run_id,
            pipeline_name,
            status="failed",
            rows_read=rows_read,
            rows_written=rows_written,
            rows_rejected=0,
            checkpoint_after={},
            request_count=0,
            error_category=type(exc).__name__,
        )
        raise


def _write_batch(database: Database, batch: list[tuple[Any, ...]]) -> int:
    with database.transaction() as connection, connection.cursor() as cursor:
        returned = execute_values(
            cursor,
            """
            INSERT INTO bronze.manual_import_rows (
                import_id, pipeline_run_id, source_name, source_row_number, row_sha256, raw_payload
            ) VALUES %s
            ON CONFLICT DO NOTHING
            RETURNING id
            """,
            batch,
            page_size=500,
            fetch=True,
        )
        return len(returned)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Validate and import a CSV/TSV/XLSX file into Bronze")
    parser.add_argument("--source", required=True)
    parser.add_argument("--file", required=True, type=Path)
    parser.add_argument("--sheet")
    parser.add_argument("--required-column", action="append", default=[])
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    warehouse = WarehouseSettings.from_environment()
    configure_logging(warehouse.log_level)
    database = Database(DatabaseSettings.from_environment(), "ela-manual-import")
    try:
        result = import_file(
            database,
            source_name=args.source,
            path=args.file,
            required_columns=tuple(args.required_column),
            sheet_name=args.sheet,
        )
        print(str(result))
        return 0
    finally:
        database.close()


if __name__ == "__main__":
    raise SystemExit(main())
