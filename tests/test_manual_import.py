from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from manual_imports.import_file import ImportValidationError, file_sha256, iter_rows


def test_csv_reader_preserves_rows(tmp_path: Path) -> None:
    path = tmp_path / "source.csv"
    path.write_text("user_id,name\n1,Asha\n2,Ravi\n", encoding="utf-8")

    headers, row_iterator = iter_rows(path)

    assert headers == ["user_id", "name"]
    assert list(row_iterator) == [
        (2, {"user_id": "1", "name": "Asha"}),
        (3, {"user_id": "2", "name": "Ravi"}),
    ]
    assert len(file_sha256(path)) == 64


def test_csv_reader_rejects_duplicate_headers(tmp_path: Path) -> None:
    path = tmp_path / "source.csv"
    path.write_text("Email,email\na,b\n", encoding="utf-8")

    with pytest.raises(ImportValidationError, match="unique"):
        iter_rows(path)


def test_xlsx_reader_uses_first_sheet(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["user_id", "name"])
    worksheet.append([1, "Asha"])
    workbook.save(path)

    headers, row_iterator = iter_rows(path)

    assert headers == ["user_id", "name"]
    assert list(row_iterator) == [(2, {"user_id": 1, "name": "Asha"})]
